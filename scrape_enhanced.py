"""
SJSU CampusConcourse — CMPE Enhanced Syllabus Scraper
=====================================================
Scrapes ALL CMPE syllabus search results, then visits each detail page
to extract the full field set matching the target xlsx schema:

    course_title, department, catalog_number, course_code,
    session, year, section, delivery, units,
    start_date, end_date, modified_date,
    college, department, description, prerequisites,
    grading_basis, source_url

Output:  cmpe_syllabi_full.xlsx   (same column layout as the sample file)

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Setup (one-time)
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    pip install playwright pandas openpyxl
    playwright install chromium

Run
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    python3 scrape_enhanced.py
"""

import re
import time
import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

BASE_URL  = "https://sjsu.campusconcourse.com"
START_URL = (f"{BASE_URL}/search?search_performed=1"
             f"&sort_by=year&descend=true&prefix=CMPE&template=non")

DETAIL_PAUSE = 0.4   # seconds between detail page fetches (be polite)


# ── helpers ────────────────────────────────────────────────────────────────────

def _excel_serial(date_str: str):
    """Convert 'MM/DD/YYYY' (or similar) to an Excel date serial number."""
    if not date_str:
        return ""
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%B %d, %Y", "%b %d, %Y"):
        try:
            d = datetime.datetime.strptime(date_str.strip(), fmt).date()
            return (d - datetime.date(1899, 12, 30)).days
        except ValueError:
            continue
    return date_str   # fall back to raw string


def _label_dd(page, label: str) -> str:
    """
    Find a <dt> whose text contains `label` (case-insensitive),
    return the text of the next sibling <dd>.
    """
    try:
        dts = page.query_selector_all("dt")
        for dt in dts:
            if label.lower() in (dt.inner_text() or "").lower():
                dd = dt.evaluate_handle(
                    "el => el.nextElementSibling"
                )
                if dd:
                    return (dd.as_element().inner_text() or "").strip()
    except Exception:
        pass
    return ""


def _first_text(page, *selectors) -> str:
    for sel in selectors:
        try:
            el = page.query_selector(sel)
            if el:
                txt = (el.inner_text() or "").strip()
                if txt:
                    return txt
        except Exception:
            pass
    return ""


# ── pass 1 : collect search-result rows ────────────────────────────────────────

def collect_search_rows(page) -> list[dict]:
    """Parse the current search-results page; return basic row dicts."""
    rows = []
    try:
        page.wait_for_selector("tbody tr", timeout=20_000)
    except PWTimeout:
        return rows

    for tr in page.query_selector_all("tbody tr"):
        # course title + URL
        a = tr.query_selector("h3 a, h4 a, h5 a, .h5 a")
        if not a:
            continue
        course_title = (a.inner_text() or "").strip()
        href = a.get_attribute("href") or ""
        source_url = href if href.startswith("http") else BASE_URL + "/" + href.lstrip("/")

        if re.search(r"template", course_title, re.I):
            continue

        cols = [c.inner_text().strip()
                for c in tr.query_selector_all("div.col-sm-3")]
        while len(cols) < 4:
            cols.append("")

        # course_code / catalog_number
        course_code = cols[0]
        catalog_number = ""
        m = re.match(r"CMPE-?\s*(\w+)", course_code, re.I)
        if m:
            catalog_number = m.group(1)
            course_code    = f"CMPE {catalog_number}"

        # session + year
        session = next(
            (s for s in ["Spring", "Fall", "Summer", "Winter"] if s in cols[1]),
            ""
        )
        m2 = re.search(r"\b(20\d{2})\b", cols[1])
        year = int(m2.group(1)) if m2 else ""

        # section number
        m3 = re.search(r"\d+", cols[2])
        section = int(m3.group()) if m3 else ""

        rows.append({
            "course_title":   course_title,
            "department":     "CMPE",
            "catalog_number": catalog_number,
            "course_code":    course_code,
            "session":        session,
            "year":           year,
            "section":        section,
            "source_url":     source_url,
        })
    return rows


def click_next(page) -> bool:
    """Submit the search form to go to the next offset page. Returns False if no Next button."""
    try:
        btn = page.query_selector("button:has-text('Next'), button:has-text('next')")
        if not btn:
            return False
        onclick = btn.get_attribute("onclick") or ""
        m = re.search(r"#offset.*?val\((\d+)\)", onclick)
        if not m:
            return False
        offset = m.group(1)
        page.evaluate(
            f"() => {{ $('#offset').val({offset}); $('#search_form').trigger('submit'); }}"
        )
        page.wait_for_selector("tbody tr", timeout=15_000)
        time.sleep(1)
        return True
    except Exception:
        return False


# ── pass 2 : scrape detail pages ───────────────────────────────────────────────

def scrape_detail(page, url: str) -> dict:
    result = {
        "delivery": "", "units": "", "start_date": "", "end_date": "",
        "modified_date": "", "college": "", "description": "",
        "prerequisites": "", "grading_basis": "",
    }
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=20_000)
        page.wait_for_load_state("networkidle", timeout=10_000)
        time.sleep(0.3)

        # delivery / instruction mode
        result["delivery"] = (
            _label_dd(page, "Instruction Mode") or
            _label_dd(page, "Delivery") or
            _label_dd(page, "Mode of Instruction")
        )

        # units
        units_raw = _label_dd(page, "Units") or _label_dd(page, "Credit Hours")
        m = re.search(r"(\d+(?:\.\d+)?)", units_raw)
        result["units"] = float(m.group(1)) if m else ""

        # dates
        start_raw    = _label_dd(page, "Start Date") or _label_dd(page, "Class Start")
        end_raw      = _label_dd(page, "End Date")   or _label_dd(page, "Class End")
        modified_raw = (
            _label_dd(page, "Modified")      or
            _label_dd(page, "Last Modified") or
            _label_dd(page, "Updated")
        )
        result["start_date"]    = _excel_serial(start_raw)
        result["end_date"]      = _excel_serial(end_raw)
        result["modified_date"] = _excel_serial(modified_raw)

        # college
        result["college"] = _label_dd(page, "College")

        # description
        result["description"] = (
            _first_text(page,
                        "#course-description",
                        ".course-description",
                        "#description",
                        "section.description p",
                        "[class*='description'] p") or
            _label_dd(page, "Description") or
            _label_dd(page, "Course Description")
        )

        # prerequisites
        result["prerequisites"] = (
            _label_dd(page, "Prerequisite") or
            _label_dd(page, "Prerequisites")
        )

        # grading
        result["grading_basis"] = (
            _label_dd(page, "Grading Basis") or
            _label_dd(page, "Grading")
        )

    except Exception as e:
        print(f" [detail error: {e}]", end="")

    return result


# ── xlsx writer ────────────────────────────────────────────────────────────────

HEADERS = [
    "course_title", "department", "catalog_number", "course_code",
    "session", "year", "section", "delivery", "units",
    "start_date", "end_date", "modified_date",
    "college", "department", "description", "prerequisites",
    "grading_basis", "source_url",
]

COL_WIDTHS = [38, 10, 14, 14, 8, 6, 8, 16, 7,
              13, 13, 14, 42, 20, 65, 55, 16, 62]

DATE_COLS  = {10, 11, 12}   # 1-indexed: start_date, end_date, modified_date
NUM_COLS   = {7, 9}          # section, units


def save_xlsx(rows: list[dict], path: str = "cmpe_syllabi_full.xlsx"):
    if not rows:
        print("No data — nothing to save.")
        return

    df = pd.DataFrame(rows)

    # Sort: newest year first → session order → course code alphabetically
    s_ord = {"Spring": 1, "Fall": 2, "Summer": 3, "Winter": 4}
    df["_s"] = df["session"].map(s_ord).fillna(5)
    df = (df.sort_values(["year", "_s", "course_code"],
                         ascending=[False, True, True])
            .drop(columns=["_s"])
            .reset_index(drop=True))

    wb = Workbook()
    ws = wb.active
    ws.title = "CMPE Syllabi"

    # ── header row ───────────────────────────────────────────────────────────
    ws.append(HEADERS)
    hdr_fill   = PatternFill("solid", start_color="1F4E79")
    hdr_font   = Font(bold=True, name="Arial", size=10, color="FFFFFF")
    hdr_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_white = Side(style="thin", color="FFFFFF")
    hdr_border = Border(left=thin_white, right=thin_white, bottom=thin_white)

    for cell in ws[1]:
        cell.font   = hdr_font
        cell.fill   = hdr_fill
        cell.alignment = hdr_align
        cell.border = hdr_border

    ws.row_dimensions[1].height = 32

    # ── data rows ────────────────────────────────────────────────────────────
    fill_even = PatternFill("solid", start_color="EBF3FB")
    fill_odd  = PatternFill("solid", start_color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    wrap      = Alignment(vertical="top", wrap_text=True)
    center    = Alignment(horizontal="center", vertical="top")

    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        fill = fill_even if r_idx % 2 == 0 else fill_odd

        # NOTE: "department" appears at column index 2 (B) AND 14 (N) in the target
        data = [
            getattr(row, "course_title",    ""),
            getattr(row, "department",      "CMPE"),
            getattr(row, "catalog_number",  ""),
            getattr(row, "course_code",     ""),
            getattr(row, "session",         ""),
            getattr(row, "year",            ""),
            getattr(row, "section",         ""),
            getattr(row, "delivery",        ""),
            getattr(row, "units",           ""),
            getattr(row, "start_date",      ""),
            getattr(row, "end_date",        ""),
            getattr(row, "modified_date",   ""),
            getattr(row, "college",         ""),
            getattr(row, "department",      "CMPE"),  # col N — matches target
            getattr(row, "description",     ""),
            getattr(row, "prerequisites",   ""),
            getattr(row, "grading_basis",   ""),
            getattr(row, "source_url",      ""),
        ]
        ws.append(data)

        for c_idx, cell in enumerate(ws[r_idx], start=1):
            cell.font = data_font
            cell.fill = fill
            if c_idx in DATE_COLS:
                cell.number_format = "MM/DD/YYYY"
                cell.alignment     = center
            elif c_idx in NUM_COLS:
                cell.alignment = center
            else:
                cell.alignment = wrap

        ws.row_dimensions[r_idx].height = 60

    # ── column widths, freeze, auto-filter ───────────────────────────────────
    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    wb.save(path)

    years = df["year"].dropna()
    print(f"\n{'─'*50}")
    print(f"  Saved : {path}")
    print(f"  Rows  : {len(df)}")
    if len(years):
        print(f"  Years : {int(years.max())} → {int(years.min())}")
    print(f"{'─'*50}\n")


# ── main ───────────────────────────────────────────────────────────────────────

def main():
    all_rows: list[dict] = []

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/124.0.0.0 Safari/537.36"
            )
        )
        page = context.new_page()

        # ── Pass 1: collect all search result rows ──────────────────────────
        print(f"\nPass 1 — collecting search results …\n")
        page.goto(START_URL, wait_until="domcontentloaded", timeout=30_000)
        time.sleep(2)

        pg = 1
        while True:
            print(f"  Page {pg} … ", end="", flush=True)
            rows = collect_search_rows(page)
            if not rows:
                print("no rows — done.")
                break
            all_rows.extend(rows)
            print(f"{len(rows)} found  (total: {len(all_rows)})")
            if not click_next(page):
                print("  No Next button — search complete.")
                break
            pg += 1

        # ── Pass 2: scrape detail pages ─────────────────────────────────────
        print(f"\nPass 2 — fetching detail pages ({len(all_rows)} syllabi) …\n")
        for i, row in enumerate(all_rows, 1):
            print(f"  [{i:>4}/{len(all_rows)}] "
                  f"{row['course_code']:12} "
                  f"{str(row.get('session','')):6} "
                  f"{str(row.get('year',''))} … ",
                  end="", flush=True)
            detail = scrape_detail(page, row["source_url"])
            row.update(detail)
            print("ok")
            time.sleep(DETAIL_PAUSE)

        browser.close()

    save_xlsx(all_rows, "cmpe_syllabi_full.xlsx")


if __name__ == "__main__":
    main()
